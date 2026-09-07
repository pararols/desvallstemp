import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm

BASE_DIR = r"c:\Users\parar\OneDrive\Documents\antigravity\Desvalls Cultura"
CONCURS_DIR = os.path.join(BASE_DIR, "concursbarrets")
WEB_DOCS_DIR = os.path.join(BASE_DIR, "web", "docs")

os.makedirs(CONCURS_DIR, exist_ok=True)
os.makedirs(WEB_DOCS_DIR, exist_ok=True)

BASES_TEXT = [
    ("1. Objecte del concurs", [
        "L'objectiu del concurs és fomentar la creativitat, la participació ciutadana i l'expressió artística en el marc de la Fira Pluja d'Art de Sant Jordi Desvalls mitjançant la decoració original de barrets."
    ]),
    ("2. Participants", [
        "• El concurs és obert a totes les persones, sense límit d'edat.",
        "• Les persones menors d'edat hi podran participar amb l'autorització dels seus pares, mares o representants legals.",
        "• Cada participant podrà presentar tants barrets com desitgi."
    ]),
    ("3. Característiques dels barrets", [
        "• Els barrets podran ser de qualsevol material, mida o estil.",
        "• La tècnica i la decoració seran totalment lliures.",
        "• Es valorarà especialment:",
        "  - l'originalitat i la creativitat;",
        "  - la qualitat artística i tècnica;",
        "  - l'acabat de la peça;",
        "  - l'impacte visual;",
        "  - l'ús de materials reciclats o sostenibles.",
        "• No s'admetran obres amb continguts ofensius, discriminatoris, violents o contraris als valors de convivència i respecte."
    ]),
    ("4. Inscripció", [
        "La participació és gratuïta.",
        "La inscripció s'haurà de formalitzar abans del dia <b>12 d'agost de 2026</b> mitjançant una de les opcions següents:",
        "• formulari habilitat per l'organització a la pàgina web: https://desvallscultura.cat/;",
        "• presencialment al Bar Social de Sant Jordi Desvalls."
    ]),
    ("5. Lliurament dels barrets", [
        "Els barrets participants s'hauran de lliurar entre el <b>31 d'agost i el 4 de setembre de 2026</b> a l'Ajuntament de Sant Jordi Desvalls, dins l'horari d'atenció al públic.",
        "En el moment del lliurament, cada participant facilitarà el seu nom i un telèfon de contacte.",
        "L'organització assignarà un número identificatiu a cada barret per garantir l'anonimat durant l'exposició i la deliberació del jurat."
    ]),
    ("6. Exposició", [
        "Tots els barrets participants s'exposaran durant la Fira Pluja d'Art de Sant Jordi Desvalls.",
        "Durant l'exposició només hi constarà el número identificatiu assignat a cada obra."
    ]),
    ("7. Jurat", [
        "El jurat serà designat per l'organització i estarà format per persones vinculades al món de l'art, la cultura, el disseny o l'associacionisme.",
        "La decisió del jurat serà motivada, definitiva i inapel·lable.",
        "En cas d'empat, la presidència del jurat disposarà de vot de qualitat."
    ]),
    ("8. Criteris de valoració", [
        "El jurat valorarà especialment:",
        "• originalitat i creativitat;",
        "• qualitat artística i tècnica;",
        "• acabat de la peça;",
        "• ús innovador o sostenible dels materials;",
        "• impacte visual del conjunt."
    ]),
    ("9. Premi", [
        "S'atorgarà un únic premi consistent en un <b>menú per a dues persones al Restaurant Els 4 Vents de Sant Jordi Desvalls</b>.",
        "El premi no podrà ser substituït pel seu valor econòmic ni bescanviat per un altre.",
        "El jurat podrà declarar el premi desert si considera que les obres presentades no reuneixen la qualitat suficient."
    ]),
    ("10. Lliurament del premi", [
        "El premi es lliurarà durant la darrera actuació de la Fira Pluja d'Art de Sant Jordi Desvalls.",
        "Si la persona guanyadora no hi és present, l'organització es posarà en contacte amb ella utilitzant les dades facilitades durant la inscripció."
    ]),
    ("11. Retorn dels barrets", [
        "Els participants podran recollir els seus barrets a l'Ajuntament durant els quinze dies naturals posteriors a la finalització de la Fira.",
        "Un cop transcorregut aquest termini sense que hagin estat retirats, l'organització podrà disposar-ne lliurement, sense dret a reclamació."
    ]),
    ("12. Responsabilitat", [
        "L'organització adoptarà les mesures raonables per garantir la correcta conservació dels barrets durant la seva custòdia i exposició.",
        "No obstant això, no respondrà dels danys ocasionats per causes de força major, robatori, actes vandàlics o altres circumstàncies alienes al seu control."
    ]),
    ("13. Drets d'imatge i difusió", [
        "Amb la participació en el concurs, els participants autoritzen l'organització a fotografiar o gravar les obres i els actes relacionats amb el concurs amb finalitats exclusivament informatives, culturals i de promoció de la Fira Pluja d'Art, a través dels seus canals habituals de comunicació i xarxes socials.",
        "Aquesta autorització no comporta cap contraprestació econòmica."
    ]),
    ("14. Protecció de dades personals", [
        "D'acord amb el Reglament (UE) 2016/679 (RGPD) i la Llei Orgànica 3/2018 (LOPDGDD), s'informa que les dades personals facilitades seran tractades per l'<b>Associació Desvalls Cultura</b> amb la finalitat de gestionar la inscripció, el desenvolupament del concurs, la comunicació amb les persones participants i el lliurament del premi.",
        "La base jurídica del tractament és el consentiment de la persona participant i l'execució de les presents bases.",
        "Les dades es conservaran durant el temps necessari per gestionar el concurs i per complir les obligacions legals que siguin aplicables.",
        "Les persones interessades poden exercir els drets d'accés, rectificació, supressió, oposició, limitació del tractament i portabilitat adreçant-se a l'<b>Associació Desvalls Cultura</b> pels canals habilitats a aquest efecte."
    ]),
    ("15. Acceptació de les bases", [
        "La participació en aquest concurs implica l'acceptació íntegra de les presents bases.",
        "L'organització es reserva la facultat d'interpretar-les i de resoldre qualsevol incidència o situació no prevista, sempre d'acord amb els principis de transparència, igualtat i bona fe."
    ])
]

# ---------------------------------------------------------
# 1. BUILD PDF FOR BASES
# ---------------------------------------------------------
def generate_pdf_bases():
    pdf_filename = os.path.join(CONCURS_DIR, "Bases_Concurs_Decoracio_Barrets_2026.pdf")
    web_pdf_filename = os.path.join(WEB_DOCS_DIR, "Bases_Concurs_Decoracio_Barrets_2026.pdf")

    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=A4,
        rightMargin=1.5*cm, leftMargin=1.5*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#581c87'),
        alignment=1,
        spaceAfter=3
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#0369a1'),
        alignment=1,
        spaceAfter=10
    )

    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor('#6d28d9'),
        spaceBefore=7,
        spaceAfter=3
    )

    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        'BulletText',
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=3
    )

    subbullet_style = ParagraphStyle(
        'SubBulletText',
        parent=body_style,
        leftIndent=24,
        firstLineIndent=-8,
        spaceAfter=2
    )

    story = []

    story.append(Paragraph("BASES DEL CONCURS DE DECORACIÓ DE BARRETS", title_style))
    story.append(Paragraph("Fira Pluja d'Art de Sant Jordi Desvalls 2026", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#6d28d9'), spaceBefore=0, spaceAfter=8))

    for title, paragraphs in BASES_TEXT:
        story.append(Paragraph(title, h2_style))
        for p in paragraphs:
            if p.startswith("  - "):
                story.append(Paragraph(p, subbullet_style))
            elif p.startswith("• "):
                story.append(Paragraph(p, bullet_style))
            else:
                story.append(Paragraph(p, body_style))

    doc.build(story)
    
    with open(pdf_filename, 'rb') as f_in:
        with open(web_pdf_filename, 'wb') as f_out:
            f_out.write(f_in.read())

    print(f"Generated PDF: {pdf_filename}")

# ---------------------------------------------------------
# 2. BUILD WORD (.DOCX) FOR BASES
# ---------------------------------------------------------
def generate_docx_bases():
    docx_filename = os.path.join(CONCURS_DIR, "Bases Concurs de Barrets 2026.docx")
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("BASES DEL CONCURS DE DECORACIÓ DE BARRETS")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(88, 28, 135)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Fira Pluja d'Art de Sant Jordi Desvalls 2026")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(3, 105, 161)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    for title, items in BASES_TEXT:
        p_h = doc.add_paragraph()
        run_h = p_h.add_run(title)
        run_h.font.name = 'Arial'
        run_h.font.size = Pt(11.5)
        run_h.font.bold = True
        run_h.font.color.rgb = RGBColor(109, 40, 217)
        p_h.paragraph_format.space_before = Pt(6)
        p_h.paragraph_format.space_after = Pt(2)

        for item in items:
            clean_item = item.replace("<b>", "").replace("</b>", "")
            p_b = doc.add_paragraph()
            run_b = p_b.add_run(clean_item)
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(10)
            p_b.paragraph_format.space_after = Pt(3)
            p_b.paragraph_format.line_spacing = 1.15

    doc.save(docx_filename)
    print(f"Generated Word Bases: {docx_filename}")

# ---------------------------------------------------------
# 3. BUILD WORD (.DOCX) FOR GRAELLA BAR SOCIAL
# ---------------------------------------------------------
def generate_docx_graella():
    docx_filename = os.path.join(CONCURS_DIR, "Graella Inscripcio Bar Social.docx")
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("FULL D'INSCRIPCIÓ PRESENCIAL - BAR SOCIAL")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(88, 28, 135)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Concurs de Decoració de Barrets · Fira Pluja d'Art 2026 (Sant Jordi Desvalls)")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(3, 105, 161)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_info = p_info.add_run(
        "📌 Data límit d'inscripció: 12 d'agost de 2026. Inscripció gratuïta i oberta a tothom.\n"
        "📦 Lliurament dels barrets: Del 31 d'agost al 4 de setembre de 2026 a l'Ajuntament de Sant Jordi Desvalls.\n"
        "🎁 Premi: Un menú per a dues persones al Restaurant Els 4 Vents."
    )
    run_info.font.name = 'Arial'
    run_info.font.size = Pt(9.5)
    run_info.font.italic = True

    table = doc.add_table(rows=16, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Nº inscripció", "Nom i Cognoms", "Telèfon", "Correu electrònic", "Edat"]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
        tcPr = hdr_cells[i]._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '6D28D9')
        tcPr.append(shd)

    for row_idx in range(1, 16):
        row_cells = table.rows[row_idx].cells
        row_cells[0].text = f"{row_idx:02d}"
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[0].paragraphs[0].runs[0].font.size = Pt(9)
        row_cells[0].paragraphs[0].runs[0].font.bold = True

    try:
        doc.save(docx_filename)
        print(f"Generated Word Graella: {docx_filename}")
    except PermissionError:
        alt_name = os.path.join(CONCURS_DIR, "Graella_Inscripcio_Bar_Social.docx")
        doc.save(alt_name)
        print(f"Generated Word Graella (nom alternatiu per permís): {alt_name}")

if __name__ == "__main__":
    generate_pdf_bases()
    generate_docx_bases()
    generate_docx_graella()
